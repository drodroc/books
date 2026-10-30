# -*- coding: utf-8 -*-
"""月度产品盈亏：导入订单和支出，把成本挂到商品上。"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import sqlite3
import threading
import uuid
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
TMP = DATA / "tmp"
DB_PATH = DATA / "ledger.db"

_lock = threading.Lock()

REFUND_DONE = "退款成功"
REFUND_CLOSED = "退款关闭"
NORMAL_STATUS = {"", "没有申请退款", REFUND_CLOSED}

# 更具体的列先匹配，避免「金额」抢走「买家实付金额」。
HEADER_FIELDS = [
    ("order_no", ("主订单编号", "订单编号", "订单号")),
    ("product_id", ("商品id", "宝贝id")),
    ("title", ("商品标题", "宝贝标题", "商品名称", "产品名称")),
    ("attr", ("商品属性", "sku属性", "销售属性")),
    ("package", ("套餐信息", "套餐")),
    ("pay_time", ("订单付款时间", "付款时间", "支付时间", "下单时间")),
    ("confirm_time", ("确认收货时间", "确认时间")),
    ("refund_status", ("退款状态",)),
    ("refund_amount", ("退款金额",)),
    ("paid", ("买家实付金额", "实付金额", "买家实付", "实收金额")),
    ("payable", ("买家应付货款", "应付金额", "买家应付")),
    ("price", ("商品价格", "单价")),
    ("qty", ("购买数量", "销售数量", "件数")),
    ("detail", ("明细", "费用明细", "费用名称", "支出项目", "费用项目")),
    ("name", ("项目", "摘要", "费用项目", "支出项目", "名称", "用途", "品名")),
    ("spent_on", ("支出日期", "发生日期", "日期")),
    ("category", ("分类", "类别", "费用类型", "科目")),
    ("note", ("备注", "说明")),
    ("link_title", ("对应商品", "关联商品", "挂靠商品", "对应产品")),
    ("amount", ("支出金额", "费用金额", "金额", "支出", "费用", "成本")),
]

FIELD_LABELS = {
    "orders": [
        ("title", "商品标题"),
        ("attr", "商品属性"),
        ("qty", "数量"),
        ("paid", "实付金额"),
        ("pay_time", "付款时间"),
        ("confirm_time", "确认收货时间"),
        ("refund_status", "退款状态"),
        ("refund_amount", "退款金额"),
        ("order_no", "订单编号"),
        ("product_id", "商品ID"),
    ],
    "aggregated": [
        ("title", "商品标题"),
        ("attr", "商品属性"),
        ("qty", "数量"),
        ("paid", "实付金额"),
        ("detail", "成本明细"),
        ("line_cost", "成本金额"),
    ],
    "expenses": [
        ("name", "项目"),
        ("amount", "金额"),
        ("spent_on", "日期"),
        ("category", "分类"),
        ("note", "备注"),
        ("link_title", "对应商品"),
    ],
}


def ensure_dirs():
    DATA.mkdir(exist_ok=True)
    TMP.mkdir(exist_ok=True)


def connect():
    ensure_dirs()
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("PRAGMA journal_mode = WAL")
    _init(conn)
    return conn


def _init(conn):
    conn.executescript(
        """
        CREATE TABLE IF NOT EXISTS imports (
            id INTEGER PRIMARY KEY,
            kind TEXT NOT NULL,
            filename TEXT,
            created_at TEXT NOT NULL,
            row_count INTEGER NOT NULL DEFAULT 0,
            month_hint TEXT,
            note TEXT
        );
        CREATE TABLE IF NOT EXISTS revenue_lines (
            id INTEGER PRIMARY KEY,
            fingerprint TEXT NOT NULL UNIQUE,
            import_id INTEGER REFERENCES imports(id) ON DELETE CASCADE,
            order_no TEXT,
            product_id TEXT,
            title TEXT NOT NULL,
            attr TEXT NOT NULL DEFAULT '',
            package TEXT,
            qty REAL NOT NULL DEFAULT 0,
            paid REAL NOT NULL DEFAULT 0,
            refund_status TEXT,
            refund_amount REAL NOT NULL DEFAULT 0,
            pay_time TEXT,
            pay_month TEXT,
            confirm_time TEXT,
            confirm_month TEXT,
            source_file TEXT
        );
        CREATE TABLE IF NOT EXISTS expenses (
            id INTEGER PRIMARY KEY,
            fingerprint TEXT UNIQUE,
            import_id INTEGER REFERENCES imports(id) ON DELETE CASCADE,
            name TEXT NOT NULL,
            amount REAL NOT NULL,
            spent_on TEXT,
            month TEXT NOT NULL,
            category TEXT DEFAULT '',
            note TEXT DEFAULT '',
            source_file TEXT
        );
        CREATE TABLE IF NOT EXISTS allocations (
            id INTEGER PRIMARY KEY,
            expense_id INTEGER NOT NULL REFERENCES expenses(id) ON DELETE CASCADE,
            title TEXT NOT NULL,
            attr TEXT NOT NULL DEFAULT '',
            amount REAL NOT NULL
        );
        CREATE TABLE IF NOT EXISTS product_meta (
            title TEXT PRIMARY KEY,
            track TEXT DEFAULT '',
            note TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS product_flags (
            title TEXT NOT NULL,
            month TEXT NOT NULL,
            no_cost INTEGER NOT NULL DEFAULT 0,
            PRIMARY KEY (title, month)
        );
        CREATE INDEX IF NOT EXISTS idx_rev_confirm ON revenue_lines(confirm_month);
        CREATE INDEX IF NOT EXISTS idx_rev_pay ON revenue_lines(pay_month);
        CREATE INDEX IF NOT EXISTS idx_rev_title ON revenue_lines(title);
        CREATE INDEX IF NOT EXISTS idx_exp_month ON expenses(month);
        """
    )
    conn.commit()


def norm_header(value):
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", "", text)
    return text.replace("：", "").replace(":", "")


def to_float(value):
    if value is None or value == "":
        return 0.0
    if isinstance(value, bool):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    text = text.replace("¥", "").replace("￥", "").replace("元", "")
    if text in ("", "无退款申请", "无", "-", "—", "none", "null"):
        return 0.0
    try:
        return float(text)
    except ValueError:
        return None


def money(value):
    return round(float(value or 0) + 1e-10, 2)


def parse_dt(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    text = str(value).strip()
    if not text or text in ("无", "-", "—"):
        return None
    text = text.replace("年", "-").replace("月", "-").replace("日", "")
    text = text.replace("/", "-").replace(".", "-")
    for fmt, size in (
        ("%Y-%m-%d %H:%M:%S", 19),
        ("%Y-%m-%d %H:%M", 16),
        ("%Y-%m-%d", 10),
        ("%Y-%m", 7),
    ):
        try:
            return datetime.strptime(text[:size], fmt)
        except ValueError:
            continue
    return None


def month_of(value):
    dt = parse_dt(value) if not isinstance(value, datetime) else value
    if isinstance(value, str) and parse_dt(value) is None:
        match = re.match(r"(\d{4})-(\d{2})", value.strip())
        if match:
            return f"{match.group(1)}-{match.group(2)}"
    if dt is None:
        return ""
    return f"{dt.year:04d}-{dt.month:02d}"


def valid_month(value):
    return bool(re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", str(value or "")))


def cell_text(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def fingerprint(parts):
    raw = "\u241f".join("" if part is None else str(part) for part in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def product_code(title):
    match = re.match(r"^([A-Za-z]{2,}(?:-[A-Za-z]*)?\d+)", title or "")
    if not match:
        return ""
    return match.group(1).upper()


def read_tabular(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return [("CSV", _read_csv(path))]
    if suffix == ".xls":
        raise ValueError("不支持老的 .xls。请用 Excel 另存为 .xlsx 再导入。")
    if suffix not in (".xlsx", ".xlsm"):
        raise ValueError("请上传 .xlsx 或 .csv 表格。")
    warnings_filter = openpyxl.styles.stylesheet
    # openpyxl 对没有默认样式的导出表会警告，不影响数据。
    import warnings

    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", category=UserWarning)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = []
        try:
            for name in wb.sheetnames:
                ws = wb[name]
                rows = []
                for index, row in enumerate(ws.iter_rows(values_only=True)):
                    if index >= 30000:
                        break
                    rows.append(list(row))
                sheets.append((name, rows))
        finally:
            wb.close()
    return sheets


def _read_csv(path: Path):
    raw = path.read_bytes()
    text = None
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ValueError("这个 CSV 的编码认不出来，请另存为 xlsx。")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;，")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    return [row for _, row in zip(range(30000), reader)]


def _header_index(rows):
    best_at = 0
    best_score = -1
    limit = min(8, len(rows))
    for index in range(limit):
        cells = [cell_text(c) for c in rows[index]]
        nonempty = [c for c in cells if c]
        if len(nonempty) < 2:
            continue
        mapping = map_headers(cells)
        score = len(mapping) * 10 + len(nonempty)
        if score > best_score:
            best_score = score
            best_at = index
    return best_at


def map_headers(headers):
    normalized = [norm_header(h) for h in headers]
    mapping = {}
    used = set()
    for field, aliases in HEADER_FIELDS:
        alias_set = {norm_header(a) for a in aliases}
        for index, header in enumerate(normalized):
            if index in used or not header:
                continue
            if header in alias_set:
                mapping[field] = index
                used.add(index)
                break
    for index, header in enumerate(normalized):
        if header in {norm_header(x) for x in ("明细", "费用明细")}:
            mapping["detail"] = index
            if index > 0 and normalized[index - 1] in {norm_header(x) for x in ("成本", "金额", "费用")}:
                mapping["line_cost"] = index - 1
    if "detail" in mapping and "line_cost" not in mapping and "amount" in mapping:
        mapping["line_cost"] = mapping["amount"]
    # 「商品标题」出现在支出表时，当作对应商品，不覆盖真正的项目列。
    if "title" in mapping and "name" in mapping and "paid" not in mapping and "attr" not in mapping:
        mapping["link_title"] = mapping.pop("title")
    return mapping


def _sheet_kind(headers, mapping):
    has = set(mapping)
    if "title" in has and ({"pay_time", "confirm_time", "order_no"} & has) and ({"paid", "qty"} & has):
        return "orders", 100 + len(has)
    if "title" in has and "attr" in has and ({"paid", "qty"} & has) and "order_no" not in has:
        return "aggregated", 80 + len(has)
    if "name" in has and "amount" in has and "attr" not in has:
        return "expenses", 70 + len(has)
    if "detail" in has and "line_cost" in has and "title" not in has:
        return "expenses", 60
    return "unknown", len(has)


def choose_sheet(sheets):
    best = None
    for name, rows in sheets:
        if not rows:
            continue
        header_at = _header_index(rows)
        headers = [cell_text(c) for c in rows[header_at]]
        mapping = map_headers(headers)
        kind, score = _sheet_kind(headers, mapping)
        # 说明页不拿来当数据。
        if norm_header(name) in ("说明", "readme"):
            score -= 1000
        item = {
            "sheet": name,
            "header_at": header_at,
            "headers": headers,
            "mapping": mapping,
            "kind": kind,
            "score": score,
            "rows": rows,
        }
        if best is None or item["score"] > best["score"]:
            best = item
    if best is None:
        raise ValueError("表格是空的。")
    return best


def _track_from_sheets(sheets):
    for name, rows in sheets:
        if "说明" not in str(name):
            continue
        for row in rows[1:12]:
            if not row:
                continue
            key = cell_text(row[0] if len(row) else "")
            if key == "赛道" and len(row) > 1:
                return cell_text(row[1])
    return ""


def _get(row, mapping, field, default=""):
    index = mapping.get(field)
    if index is None or index >= len(row):
        return default
    return row[index]


def classify_refund(status):
    status = cell_text(status)
    if status == REFUND_DONE:
        return "refund"
    if status in NORMAL_STATUS:
        return "normal"
    if "退款" in status or "退货" in status:
        return "pending"
    return "normal"


def _data_rows(chosen):
    rows = []
    last_title = ""
    for raw in chosen["rows"][chosen["header_at"] + 1 :]:
        if not raw or all(cell_text(c) == "" for c in raw):
            continue
        title = cell_text(_get(raw, chosen["mapping"], "title", ""))
        if title in ("合计", "总计") or title.startswith("合计"):
            continue
        if title:
            last_title = title
        elif last_title and chosen["kind"] == "aggregated":
            title = last_title
        rows.append((raw, title))
    return rows


def analyze_file(path: Path):
    sheets = read_tabular(path)
    chosen = choose_sheet(sheets)
    track = _track_from_sheets(sheets)
    kind = chosen["kind"]
    mapping = chosen["mapping"]
    warnings = []
    months = {}
    cost_rows = 0
    usable = 0
    preview = []
    if kind == "unknown":
        warnings.append("没有认出这是收益表还是支出表，请在下面指定每一列是什么。")
    for raw, title in _data_rows(chosen):
        if kind == "orders":
            paid = to_float(_get(raw, mapping, "paid"))
            qty = to_float(_get(raw, mapping, "qty"))
            if not title or paid is None or qty is None:
                continue
            if paid == 0 and qty == 0:
                continue
            usable += 1
            pay_m = month_of(_get(raw, mapping, "pay_time"))
            confirm_m = month_of(_get(raw, mapping, "confirm_time"))
            key = confirm_m or pay_m or "（没有日期）"
            months[key] = months.get(key, 0) + 1
            if len(preview) < 6:
                preview.append(
                    {
                        "商品": title,
                        "属性": cell_text(_get(raw, mapping, "attr")),
                        "数量": qty,
                        "实付": paid,
                        "付款": cell_text(_get(raw, mapping, "pay_time")),
                        "确认收货": cell_text(_get(raw, mapping, "confirm_time")) or "无",
                        "退款": cell_text(_get(raw, mapping, "refund_status")) or "无",
                    }
                )
        elif kind == "aggregated":
            paid = to_float(_get(raw, mapping, "paid"))
            qty = to_float(_get(raw, mapping, "qty"))
            if title and paid is not None and qty is not None and not (paid == 0 and qty == 0):
                usable += 1
                if len(preview) < 6:
                    preview.append(
                        {
                            "商品": title,
                            "属性": cell_text(_get(raw, mapping, "attr")),
                            "数量": qty,
                            "实付": paid,
                        }
                    )
            detail = cell_text(_get(raw, mapping, "detail"))
            cost = to_float(_get(raw, mapping, "line_cost")) if "line_cost" in mapping else None
            if detail and cost not in (None, 0):
                cost_rows += 1
        elif kind == "expenses":
            name = cell_text(_get(raw, mapping, "name")) or cell_text(_get(raw, mapping, "detail"))
            amount = to_float(_get(raw, mapping, "amount"))
            if "amount" not in mapping and "line_cost" in mapping:
                amount = to_float(_get(raw, mapping, "line_cost"))
            if not name or amount is None or amount == 0:
                continue
            usable += 1
            spent = month_of(_get(raw, mapping, "spent_on"))
            if spent:
                months[spent] = months.get(spent, 0) + 1
            if len(preview) < 8:
                preview.append(
                    {
                        "项目": name,
                        "金额": amount,
                        "日期": cell_text(_get(raw, mapping, "spent_on")),
                        "分类": cell_text(_get(raw, mapping, "category")),
                        "对应商品": cell_text(_get(raw, mapping, "link_title")) or title,
                    }
                )
    needs = False
    if kind == "orders" and ("title" not in mapping or "paid" not in mapping):
        needs = True
    if kind == "aggregated" and ("title" not in mapping or "paid" not in mapping):
        needs = True
    if kind == "expenses" and ("name" not in mapping or ("amount" not in mapping and "line_cost" not in mapping)):
        needs = True
    if kind == "unknown":
        needs = True
    if kind == "orders" and "confirm_time" not in mapping and "pay_time" not in mapping:
        warnings.append("没有付款时间，也没有确认收货时间。导入时要指定一个月份，这些行会全部算成该月的确认收货。")
    if kind == "aggregated":
        warnings.append("这是已经合并过的商品表，里面没有每一笔订单的确认收货时间。导入时要指定月份，金额全部记为该月确认收货。")
        if cost_rows:
            warnings.append(f"表里还有 {cost_rows} 笔写在「明细」旁边的成本。可以一并导入，并自动挂到对应商品上。")
        warnings.append("如果这个月的订单明细已经导入过，不要再把合并表当收益导入，否则收入会加两次。")
    undated = months.get("（没有日期）", 0)
    if kind == "orders" and undated:
        warnings.append(
            f"有 {undated} 行没有付款时间，也没有确认收货时间，多半是还没付款的记录。这些行默认不导入。"
        )
    suggested = ""
    real_months = [m for m in months if valid_month(m)]
    if real_months:
        suggested = sorted(real_months)[-1]
    return {
        "kind": kind,
        "sheet": chosen["sheet"],
        "headers": [h for h in chosen["headers"] if h] or chosen["headers"],
        "header_row": chosen["headers"],
        "mapping": mapping,
        "needs_mapping": needs,
        "fields": FIELD_LABELS.get(kind if kind != "unknown" else "orders", FIELD_LABELS["orders"]),
        "preview": preview,
        "stats": {
            "rows": usable,
            "months": months,
            "cost_rows": cost_rows,
            "track": track,
        },
        "suggested_month": suggested,
        "warnings": warnings,
    }


def save_upload(filename: str, data: bytes):
    ensure_dirs()
    token = uuid.uuid4().hex
    suffix = Path(filename or "upload.xlsx").suffix.lower()
    if suffix not in (".xlsx", ".xlsm", ".csv", ".xls"):
        suffix = ".xlsx"
    path = TMP / f"{token}{suffix}"
    path.write_bytes(data)
    meta = {"filename": Path(filename or path.name).name, "path": str(path)}
    (TMP / f"{token}.json").write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    return token, path


def load_upload(token: str):
    if not re.fullmatch(r"[0-9a-f]{32}", token or ""):
        raise ValueError("导入凭证无效，请重新选择文件。")
    meta_path = TMP / f"{token}.json"
    if not meta_path.exists():
        raise ValueError("找不到刚才上传的文件，请重新选择。")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    path = Path(meta["path"])
    if not path.exists():
        raise ValueError("上传的文件已经丢失，请重新选择。")
    return meta, path


def _apply_mapping(chosen, mapping):
    if mapping:
        cleaned = {}
        for key, value in mapping.items():
            if value is None or value == "":
                continue
            cleaned[key] = int(value)
        chosen["mapping"] = cleaned
    kind, _ = _sheet_kind(chosen["headers"], chosen["mapping"])
    # 用户指定的类型优先。调用方会覆盖 kind。
    chosen["detected_kind"] = kind
    return chosen


def commit_import(conn, token, kind, month, include_costs, mapping, track, allow_overlap=False, assign_undated=False):
    meta, path = load_upload(token)
    sheets = read_tabular(path)
    chosen = choose_sheet(sheets)
    file_track = track or _track_from_sheets(sheets)
    if mapping:
        chosen["mapping"] = {k: int(v) for k, v in mapping.items() if v is not None and v != ""}
    if kind not in ("orders", "aggregated", "expenses"):
        kind = chosen["kind"]
    if kind == "unknown":
        raise ValueError("还没认出表格类型，请先指定列。")
    chosen["kind"] = kind
    if kind in ("aggregated", "expenses") and not valid_month(month):
        # 支出如果每行自己有日期，月份可以为空，空的行再用 month。
        if kind == "aggregated":
            raise ValueError("合并表要选择它代表的月份。")
    if kind == "orders" and "confirm_time" not in chosen["mapping"] and "pay_time" not in chosen["mapping"]:
        if not valid_month(month):
            raise ValueError("这张表没有日期，请选择归入哪个月。")
    if assign_undated and not valid_month(month):
        raise ValueError("要把没有日期的行归进去，请先选月份。")
    overlap_orders = 0
    if kind == "aggregated" and valid_month(month):
        overlap_orders = conn.execute(
            """SELECT COUNT(*) c FROM revenue_lines
               WHERE order_no != '' AND (confirm_month = ? OR pay_month = ?)""",
            (month, month),
        ).fetchone()["c"]
        if overlap_orders and not allow_overlap:
            raise ValueError(
                "这个月已经有订单明细。再导入合并表会把收入加一次。"
                "如果这份合并表才是这个月的收益来源，请勾选「仍然导入」。"
            )
    now = datetime.now().isoformat(timespec="seconds")
    cur = conn.execute(
        "INSERT INTO imports (kind, filename, created_at, row_count, month_hint, note) VALUES (?,?,?,?,?,?)",
        (kind, meta["filename"], now, 0, month or "", file_track or ""),
    )
    import_id = cur.lastrowid
    inserted = 0
    updated = 0
    skipped = 0
    costs = 0
    linked = 0
    titles = set()
    if kind == "orders":
        inserted, updated, skipped, titles = _commit_orders(
            conn, chosen, import_id, meta["filename"], month, assign_undated
        )
    elif kind == "aggregated":
        inserted, skipped, titles = _commit_aggregated(conn, chosen, import_id, meta["filename"], month)
        if include_costs:
            costs, linked = _commit_embedded_costs(conn, chosen, import_id, meta["filename"], month)
    else:
        inserted, skipped, linked = _commit_expenses(conn, chosen, import_id, meta["filename"], month)
    if file_track:
        for title in titles:
            conn.execute(
                "INSERT INTO product_meta (title, track, note) VALUES (?,?,?) "
                "ON CONFLICT(title) DO UPDATE SET track = CASE WHEN product_meta.track = '' THEN excluded.track ELSE product_meta.track END",
                (title, file_track, ""),
            )
    conn.execute("UPDATE imports SET row_count = ? WHERE id = ?", (inserted + costs, import_id))
    conn.commit()
    try:
        path.unlink()
        (TMP / f"{token}.json").unlink()
    except OSError:
        pass
    return {
        "import_id": import_id,
        "kind": kind,
        "inserted": inserted,
        "updated": updated,
        "skipped": skipped,
        "costs": costs,
        "linked": linked,
        "month": month,
        "overlap_orders": overlap_orders,
    }


def _commit_orders(conn, chosen, import_id, filename, forced_month, assign_undated=False):
    mapping = chosen["mapping"]
    inserted = updated = skipped = 0
    titles = set()
    seen = set()
    for raw, title in _data_rows(chosen):
        if not title:
            skipped += 1
            continue
        qty = to_float(_get(raw, mapping, "qty")) or 0
        paid = to_float(_get(raw, mapping, "paid"))
        if paid is None:
            skipped += 1
            continue
        if qty == 0 and paid == 0:
            skipped += 1
            continue
        attr = cell_text(_get(raw, mapping, "attr"))
        order_no = cell_text(_get(raw, mapping, "order_no"))
        product_id = cell_text(_get(raw, mapping, "product_id"))
        pay_time = cell_text(_get(raw, mapping, "pay_time"))
        confirm_time = cell_text(_get(raw, mapping, "confirm_time"))
        pay_month = month_of(pay_time)
        confirm_month = month_of(confirm_time)
        if not pay_month and not confirm_month:
            # 整张表都没有日期列时，用用户选的月份。否则没有日期的行默认跳过。
            sheet_has_dates = "pay_time" in mapping or "confirm_time" in mapping
            if (not sheet_has_dates or assign_undated) and valid_month(forced_month):
                confirm_month = forced_month
                confirm_time = f"{forced_month}-01"
            else:
                skipped += 1
                continue
        status = cell_text(_get(raw, mapping, "refund_status"))
        refund_amount = to_float(_get(raw, mapping, "refund_amount")) or 0
        base = fingerprint([order_no, product_id, title, attr, pay_time, money(qty), money(paid)])
        fp = base
        bump = 1
        while fp in seen:
            bump += 1
            fp = fingerprint([base, bump])
        seen.add(fp)
        existing = conn.execute("SELECT id, import_id FROM revenue_lines WHERE fingerprint = ?", (fp,)).fetchone()
        if existing:
            conn.execute(
                """UPDATE revenue_lines
                   SET refund_status=?, refund_amount=?, pay_time=?, pay_month=?,
                       confirm_time=?, confirm_month=?, qty=?, paid=?, import_id=?, source_file=?
                   WHERE id=?""",
                (
                    status, money(refund_amount), pay_time, pay_month, confirm_time, confirm_month,
                    qty, money(paid), import_id, filename, existing["id"],
                ),
            )
            updated += 1
        else:
            conn.execute(
                """INSERT INTO revenue_lines (
                    fingerprint, import_id, order_no, product_id, title, attr, package,
                    qty, paid, refund_status, refund_amount, pay_time, pay_month,
                    confirm_time, confirm_month, source_file
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (
                    fp, import_id, order_no, product_id, title, attr,
                    cell_text(_get(raw, mapping, "package")),
                    qty, money(paid), status, money(refund_amount),
                    pay_time, pay_month, confirm_time, confirm_month, filename,
                ),
            )
            inserted += 1
        titles.add(title)
    return inserted, updated, skipped, titles


def _commit_aggregated(conn, chosen, import_id, filename, month):
    mapping = chosen["mapping"]
    inserted = skipped = 0
    titles = set()
    for raw, title in _data_rows(chosen):
        qty = to_float(_get(raw, mapping, "qty")) or 0
        paid = to_float(_get(raw, mapping, "paid"))
        if not title or paid is None or (qty == 0 and paid == 0):
            continue
        attr = cell_text(_get(raw, mapping, "attr"))
        fp = fingerprint(["agg", month, title, attr])
        existing = conn.execute("SELECT id FROM revenue_lines WHERE fingerprint = ?", (fp,)).fetchone()
        if existing:
            skipped += 1
            continue
        conn.execute(
            """INSERT INTO revenue_lines (
                fingerprint, import_id, order_no, product_id, title, attr, package,
                qty, paid, refund_status, refund_amount, pay_time, pay_month,
                confirm_time, confirm_month, source_file
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                fp, import_id, "", "", title, attr, "",
                qty, money(paid), "没有申请退款", 0,
                f"{month}-01", month, f"{month}-01", month, filename,
            ),
        )
        inserted += 1
        titles.add(title)
    return inserted, skipped, titles


def _commit_embedded_costs(conn, chosen, import_id, filename, month):
    mapping = chosen["mapping"]
    if "detail" not in mapping or "line_cost" not in mapping:
        return 0, 0
    inserted = linked = 0
    last_title = ""
    for index, (raw, title) in enumerate(_data_rows(chosen)):
        if title:
            last_title = title
        detail = cell_text(_get(raw, mapping, "detail"))
        cost = to_float(_get(raw, mapping, "line_cost"))
        if not detail or cost is None or cost == 0 or not last_title:
            continue
        if cost < 0:
            continue
        fp = fingerprint(["cost", month, last_title, detail, money(cost), index])
        if conn.execute("SELECT 1 FROM expenses WHERE fingerprint = ?", (fp,)).fetchone():
            continue
        cur = conn.execute(
            """INSERT INTO expenses (fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (fp, import_id, detail, money(cost), "", month, "", f"来自合并表，挂在 {last_title}", filename),
        )
        conn.execute(
            "INSERT INTO allocations (expense_id, title, attr, amount) VALUES (?,?,?,?)",
            (cur.lastrowid, last_title, "", money(cost)),
        )
        inserted += 1
        linked += 1
    return inserted, linked


def _commit_expenses(conn, chosen, import_id, filename, month):
    mapping = chosen["mapping"]
    inserted = skipped = linked = 0
    known = [r[0] for r in conn.execute("SELECT DISTINCT title FROM revenue_lines")]
    for index, (raw, title) in enumerate(_data_rows(chosen)):
        name = cell_text(_get(raw, mapping, "name")) or cell_text(_get(raw, mapping, "detail"))
        if "amount" in mapping:
            amount = to_float(_get(raw, mapping, "amount"))
        else:
            amount = to_float(_get(raw, mapping, "line_cost"))
        if not name or amount is None or amount <= 0:
            skipped += 1
            continue
        spent_on = cell_text(_get(raw, mapping, "spent_on"))
        row_month = month_of(spent_on) or (month if valid_month(month) else "")
        if not valid_month(row_month):
            raise ValueError(f"支出「{name}」没有日期，请先选择这些支出算在哪个月。")
        link = cell_text(_get(raw, mapping, "link_title")) or title
        category = cell_text(_get(raw, mapping, "category"))
        note = cell_text(_get(raw, mapping, "note"))
        fp = fingerprint(["exp", filename, index, name, money(amount), spent_on, link])
        if conn.execute("SELECT 1 FROM expenses WHERE fingerprint = ?", (fp,)).fetchone():
            skipped += 1
            continue
        cur = conn.execute(
            """INSERT INTO expenses (fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
               VALUES (?,?,?,?,?,?,?,?,?)""",
            (fp, import_id, name, money(amount), spent_on, row_month, category, note, filename),
        )
        inserted += 1
        matched = _match_title(link, known)
        if matched:
            conn.execute(
                "INSERT INTO allocations (expense_id, title, attr, amount) VALUES (?,?,?,?)",
                (cur.lastrowid, matched, "", money(amount)),
            )
            linked += 1
    return inserted, skipped, linked


def _match_title(link, titles):
    link = (link or "").strip()
    if not link:
        return ""
    if link in titles:
        return link
    code = product_code(link)
    if code:
        hits = [t for t in titles if product_code(t) == code or code in t.upper()]
        if len(hits) == 1:
            return hits[0]
    if len(link) >= 4:
        hits = [t for t in titles if link in t or t in link]
        if len(hits) == 1:
            return hits[0]
    return ""


def _suggest(name, titles):
    code_in_name = []
    upper = (name or "").upper()
    for title in titles:
        code = product_code(title)
        if code and code in upper:
            code_in_name.append(title)
    if len(code_in_name) == 1:
        return code_in_name[0]
    return ""


def add_expense(conn, name, amount, month, category="", note="", spent_on="", title=""):
    name = (name or "").strip()
    if not name:
        raise ValueError("请填写支出项目。")
    amount = to_float(amount)
    if amount is None or amount <= 0:
        raise ValueError("金额要大于 0。")
    if not valid_month(month):
        raise ValueError("请选择这笔支出属于哪个月。")
    cur = conn.execute(
        """INSERT INTO expenses (fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (None, None, name, money(amount), spent_on or "", month, category or "", note or "", "手工"),
    )
    expense_id = cur.lastrowid
    if title:
        conn.execute(
            "INSERT INTO allocations (expense_id, title, attr, amount) VALUES (?,?,?,?)",
            (expense_id, title, "", money(amount)),
        )
    conn.commit()
    return expense_id


def update_expense(conn, expense_id, fields):
    row = conn.execute("SELECT * FROM expenses WHERE id = ?", (expense_id,)).fetchone()
    if not row:
        raise ValueError("找不到这笔支出。")
    name = fields.get("name", row["name"]).strip()
    amount = money(fields.get("amount", row["amount"]))
    month = fields.get("month", row["month"])
    if not name:
        raise ValueError("项目不能是空的。")
    if amount <= 0:
        raise ValueError("金额要大于 0。")
    if not valid_month(month):
        raise ValueError("月份不对。")
    allocated = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM allocations WHERE expense_id = ?", (expense_id,)).fetchone()["s"]
    if amount + 0.001 < allocated:
        raise ValueError("金额不能小于已经挂到商品上的合计。先把多挂的部分拿下来。")
    conn.execute(
        "UPDATE expenses SET name=?, amount=?, month=?, category=?, note=?, spent_on=? WHERE id=?",
        (
            name,
            amount,
            month,
            fields.get("category", row["category"]) or "",
            fields.get("note", row["note"]) or "",
            fields.get("spent_on", row["spent_on"]) or "",
            expense_id,
        ),
    )
    conn.commit()


def delete_expense(conn, expense_id):
    conn.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
    conn.commit()


def allocate(conn, expense_id, title, amount=None, attr=""):
    title = (title or "").strip()
    attr = attr or ""
    if not title:
        raise ValueError("要挂到哪个商品？")
    expense = conn.execute("SELECT * FROM expenses WHERE id = ?", (expense_id,)).fetchone()
    if not expense:
        raise ValueError("找不到这笔支出。")
    used = conn.execute("SELECT COALESCE(SUM(amount),0) s FROM allocations WHERE expense_id = ?", (expense_id,)).fetchone()["s"]
    remaining = money(expense["amount"] - used)
    if amount is None:
        amount = remaining
    amount = money(amount)
    if amount <= 0:
        raise ValueError("没有剩下的金额可以挂了。")
    if amount - remaining > 0.009:
        raise ValueError(f"这笔支出还剩 {remaining:.2f}，不够挂出 {amount:.2f}。")
    existing = conn.execute(
        "SELECT id, amount FROM allocations WHERE expense_id=? AND title=? AND attr=?",
        (expense_id, title, attr),
    ).fetchone()
    if existing:
        conn.execute("UPDATE allocations SET amount=? WHERE id=?", (money(existing["amount"] + amount), existing["id"]))
    else:
        conn.execute(
            "INSERT INTO allocations (expense_id, title, attr, amount) VALUES (?,?,?,?)",
            (expense_id, title, attr, amount),
        )
    conn.execute(
        "UPDATE product_flags SET no_cost=0 WHERE title=? AND month=?",
        (title, expense["month"]),
    )
    conn.commit()


def update_allocation(conn, allocation_id, amount=None, title=None, attr=None):
    row = conn.execute(
        """SELECT a.*, e.amount AS expense_amount, e.month
           FROM allocations a JOIN expenses e ON e.id = a.expense_id
           WHERE a.id = ?""",
        (allocation_id,),
    ).fetchone()
    if not row:
        raise ValueError("找不到这条分摊。")
    new_title = row["title"] if title is None else title.strip()
    new_attr = row["attr"] if attr is None else attr
    new_amount = row["amount"] if amount is None else money(amount)
    if not new_title:
        raise ValueError("商品不能是空的。")
    if new_amount <= 0:
        raise ValueError("金额要大于 0。可以改成拿下来。")
    others = conn.execute(
        "SELECT COALESCE(SUM(amount),0) s FROM allocations WHERE expense_id=? AND id<>?",
        (row["expense_id"], allocation_id),
    ).fetchone()["s"]
    if new_amount - money(row["expense_amount"] - others) > 0.009:
        raise ValueError("超过了这笔支出的金额。")
    twin = conn.execute(
        "SELECT id, amount FROM allocations WHERE expense_id=? AND title=? AND attr=? AND id<>?",
        (row["expense_id"], new_title, new_attr, allocation_id),
    ).fetchone()
    if twin:
        conn.execute("UPDATE allocations SET amount=? WHERE id=?", (money(twin["amount"] + new_amount), twin["id"]))
        conn.execute("DELETE FROM allocations WHERE id=?", (allocation_id,))
    else:
        conn.execute(
            "UPDATE allocations SET amount=?, title=?, attr=? WHERE id=?",
            (new_amount, new_title, new_attr, allocation_id),
        )
    conn.commit()


def delete_allocation(conn, allocation_id):
    conn.execute("DELETE FROM allocations WHERE id=?", (allocation_id,))
    conn.commit()


def set_product_meta(conn, title, track=None, note=None):
    current = conn.execute("SELECT * FROM product_meta WHERE title=?", (title,)).fetchone()
    new_track = current["track"] if current and track is None else (track or "")
    new_note = current["note"] if current and note is None else (note or "")
    conn.execute(
        "INSERT INTO product_meta (title, track, note) VALUES (?,?,?) "
        "ON CONFLICT(title) DO UPDATE SET track=excluded.track, note=excluded.note",
        (title, new_track, new_note),
    )
    conn.commit()


def set_no_cost(conn, title, month, no_cost):
    if not valid_month(month):
        raise ValueError("月份不对。")
    conn.execute(
        "INSERT INTO product_flags (title, month, no_cost) VALUES (?,?,?) "
        "ON CONFLICT(title, month) DO UPDATE SET no_cost=excluded.no_cost",
        (title, month, 1 if no_cost else 0),
    )
    conn.commit()


def delete_import(conn, import_id):
    conn.execute("DELETE FROM imports WHERE id=?", (import_id,))
    conn.commit()


def list_months(conn):
    found = set()
    for sql in (
        "SELECT pay_month m FROM revenue_lines WHERE pay_month != ''",
        "SELECT confirm_month m FROM revenue_lines WHERE confirm_month != ''",
        "SELECT month m FROM expenses",
    ):
        for row in conn.execute(sql):
            if valid_month(row["m"]):
                found.add(row["m"])
    return sorted(found)


def build_state(conn, month, basis="confirmed"):
    months = list_months(conn)
    if not valid_month(month):
        month = months[-1] if months else datetime.now().strftime("%Y-%m")
    if basis not in ("confirmed", "all"):
        basis = "confirmed"
    products = {}

    def bucket(title):
        if title not in products:
            products[title] = {
                "title": title,
                "track": "",
                "confirmed_qty": 0.0,
                "confirmed_amt": 0.0,
                "unconfirmed_qty": 0.0,
                "unconfirmed_amt": 0.0,
                "refund_qty": 0.0,
                "refund_amt": 0.0,
                "pending_count": 0,
                "pending_amt": 0.0,
                "attrs": {},
                "allocations": [],
                "no_cost": False,
            }
        return products[title]

    def attr_bucket(prod, attr):
        attrs = prod["attrs"]
        if attr not in attrs:
            attrs[attr] = {
                "attr": attr or "（无属性）",
                "attr_raw": attr or "",
                "confirmed_qty": 0.0,
                "confirmed_amt": 0.0,
                "unconfirmed_qty": 0.0,
                "unconfirmed_amt": 0.0,
            }
        return attrs[attr]

    rows = conn.execute(
        """SELECT title, attr, qty, paid, refund_status, refund_amount, pay_month, confirm_month
           FROM revenue_lines
           WHERE confirm_month = ? OR pay_month = ?""",
        (month, month),
    )
    for row in rows:
        kind = classify_refund(row["refund_status"])
        prod = bucket(row["title"])
        qty = row["qty"] or 0
        paid = row["paid"] or 0
        if kind == "refund":
            in_month = row["confirm_month"] == month or (not row["confirm_month"] and row["pay_month"] == month)
            if in_month:
                refund_amt = row["refund_amount"] if (row["refund_amount"] or 0) > 0 else paid
                prod["refund_qty"] += qty
                prod["refund_amt"] += refund_amt
            continue
        attr = attr_bucket(prod, row["attr"] or "")
        if row["confirm_month"] == month:
            prod["confirmed_qty"] += qty
            prod["confirmed_amt"] += paid
            attr["confirmed_qty"] += qty
            attr["confirmed_amt"] += paid
            if kind == "pending":
                prod["pending_count"] += 1
                prod["pending_amt"] += paid
        elif row["pay_month"] == month and not row["confirm_month"]:
            prod["unconfirmed_qty"] += qty
            prod["unconfirmed_amt"] += paid
            attr["unconfirmed_qty"] += qty
            attr["unconfirmed_amt"] += paid
            if kind == "pending":
                prod["pending_count"] += 1
                prod["pending_amt"] += paid

    backlog = conn.execute(
        """SELECT COALESCE(SUM(qty),0) qty, COALESCE(SUM(paid),0) amt
           FROM revenue_lines
           WHERE pay_month != '' AND pay_month < ?
             AND (confirm_month IS NULL OR confirm_month = '')
             AND refund_status != ?""",
        (month, REFUND_DONE),
    ).fetchone()

    meta = {r["title"]: r for r in conn.execute("SELECT title, track, note FROM product_meta")}
    flags = {
        r["title"]: r["no_cost"]
        for r in conn.execute("SELECT title, no_cost FROM product_flags WHERE month=?", (month,))
    }
    for title, prod in products.items():
        if title in meta:
            prod["track"] = meta[title]["track"] or ""
        prod["no_cost"] = bool(flags.get(title))

    expenses = []
    for row in conn.execute("SELECT * FROM expenses WHERE month=? ORDER BY id", (month,)):
        allocs = conn.execute(
            "SELECT id, title, attr, amount FROM allocations WHERE expense_id=? ORDER BY id",
            (row["id"],),
        ).fetchall()
        used = money(sum(a["amount"] for a in allocs))
        item = {
            "id": row["id"],
            "name": row["name"],
            "amount": money(row["amount"]),
            "remaining": money(row["amount"] - used),
            "spent_on": row["spent_on"] or "",
            "month": row["month"],
            "category": row["category"] or "",
            "note": row["note"] or "",
            "allocations": [
                {"id": a["id"], "title": a["title"], "attr": a["attr"] or "", "amount": money(a["amount"])}
                for a in allocs
            ],
        }
        expenses.append(item)
        for alloc in item["allocations"]:
            prod = bucket(alloc["title"])
            prod["allocations"].append(
                {
                    "id": alloc["id"],
                    "expense_id": row["id"],
                    "name": row["name"],
                    "amount": alloc["amount"],
                    "attr": alloc["attr"],
                    "category": item["category"],
                }
            )

    titles = list(products)
    for item in expenses:
        item["suggestion"] = _suggest(item["name"], titles) if item["remaining"] > 0 else ""

    product_list = []
    profit_count = loss_count = even_count = unknown_count = 0
    for prod in products.values():
        prod["confirmed_qty"] = money(prod["confirmed_qty"])
        prod["confirmed_amt"] = money(prod["confirmed_amt"])
        prod["unconfirmed_qty"] = money(prod["unconfirmed_qty"])
        prod["unconfirmed_amt"] = money(prod["unconfirmed_amt"])
        prod["refund_qty"] = money(prod["refund_qty"])
        prod["refund_amt"] = money(prod["refund_amt"])
        prod["pending_amt"] = money(prod["pending_amt"])
        cost = money(sum(a["amount"] for a in prod["allocations"]))
        prod["cost"] = cost
        revenue = prod["confirmed_amt"] + (prod["unconfirmed_amt"] if basis == "all" else 0)
        known = cost > 0 or prod["no_cost"]
        if revenue == 0 and cost == 0:
            prod["profit"] = 0
            prod["status"] = "none"
        elif not known:
            prod["profit"] = None
            prod["status"] = "unknown"
            unknown_count += 1
        else:
            prod["profit"] = money(revenue - cost)
            if prod["profit"] > 0.004:
                prod["status"] = "profit"
                profit_count += 1
            elif prod["profit"] < -0.004:
                prod["status"] = "loss"
                loss_count += 1
            else:
                prod["status"] = "even"
                even_count += 1
        attr_rows = []
        for attr in prod["attrs"].values():
            attr["confirmed_amt"] = money(attr["confirmed_amt"])
            attr["confirmed_qty"] = money(attr["confirmed_qty"])
            attr["unconfirmed_amt"] = money(attr["unconfirmed_amt"])
            attr["unconfirmed_qty"] = money(attr["unconfirmed_qty"])
            base_amt = attr["confirmed_amt"] + (attr["unconfirmed_amt"] if basis == "all" else 0)
            share = (base_amt / revenue) if revenue else 0
            attr["share"] = round(share, 6)
            direct = money(sum(a["amount"] for a in prod["allocations"] if a["attr"] == attr["attr_raw"] and attr["attr_raw"]))
            shared = money(sum(a["amount"] for a in prod["allocations"] if not a["attr"]))
            attr["direct_cost"] = direct
            attr["estimated_cost"] = money(direct + shared * share) if known else None
            attr["estimated_profit"] = money(base_amt - attr["estimated_cost"]) if known and attr["estimated_cost"] is not None else None
            attr_rows.append(attr)
        attr_rows.sort(key=lambda a: (a["confirmed_amt"] + a["unconfirmed_amt"]), reverse=True)
        prod["attrs"] = attr_rows
        prod["attr_count"] = len(attr_rows)
        if prod["confirmed_amt"] or prod["unconfirmed_amt"] or prod["allocations"] or prod["refund_amt"]:
            product_list.append(prod)
    product_list.sort(key=lambda p: p["confirmed_amt"], reverse=True)

    confirmed_amt = money(sum(p["confirmed_amt"] for p in product_list))
    confirmed_qty = money(sum(p["confirmed_qty"] for p in product_list))
    unconfirmed_amt = money(sum(p["unconfirmed_amt"] for p in product_list))
    unconfirmed_qty = money(sum(p["unconfirmed_qty"] for p in product_list))
    refund_amt = money(sum(p["refund_amt"] for p in product_list))
    refund_qty = money(sum(p["refund_qty"] for p in product_list))
    pending_count = sum(p["pending_count"] for p in product_list)
    pending_amt = money(sum(p["pending_amt"] for p in product_list))
    expense_total = money(sum(e["amount"] for e in expenses))
    allocated_total = money(sum(e["amount"] - e["remaining"] for e in expenses))
    unallocated_total = money(expense_total - allocated_total)
    revenue_for_profit = confirmed_amt + (unconfirmed_amt if basis == "all" else 0)
    has_expenses = bool(expenses)
    month_profit = money(revenue_for_profit - expense_total) if has_expenses else None

    summary = {
        "month": month,
        "basis": basis,
        "confirmed_qty": confirmed_qty,
        "confirmed_amt": confirmed_amt,
        "unconfirmed_qty": unconfirmed_qty,
        "unconfirmed_amt": unconfirmed_amt,
        "refund_qty": refund_qty,
        "refund_amt": refund_amt,
        "pending_count": pending_count,
        "pending_amt": pending_amt,
        "backlog_qty": money(backlog["qty"] or 0),
        "backlog_amt": money(backlog["amt"] or 0),
        "expense_total": expense_total,
        "allocated_total": allocated_total,
        "unallocated_total": unallocated_total,
        "revenue_for_profit": revenue_for_profit,
        "month_profit": month_profit,
        "has_expenses": has_expenses,
        "product_count": len(product_list),
        "profit_count": profit_count,
        "loss_count": loss_count,
        "even_count": even_count,
        "unknown_count": unknown_count,
    }
    summary["sentence"] = conclusion_sentence(summary)
    imports = [
        dict(r)
        for r in conn.execute("SELECT id, kind, filename, created_at, row_count, month_hint, note FROM imports ORDER BY id DESC")
    ]
    tracks = sorted({p["track"] for p in product_list if p["track"]})
    categories = sorted({e["category"] for e in expenses if e["category"]})
    return {
        "month": month,
        "months": months,
        "basis": basis,
        "summary": summary,
        "products": product_list,
        "expenses": expenses,
        "imports": imports,
        "tracks": tracks,
        "categories": categories,
    }


def conclusion_sentence(summary):
    month = summary["month"]
    revenue = summary["revenue_for_profit"]
    basis_text = "确认收货" if summary["basis"] == "confirmed" else "确认收货加未确认"
    if not summary["has_expenses"] and revenue == 0 and summary["product_count"] == 0:
        return "还没有数据。先导入收益表，再导入支出表。"
    if not summary["has_expenses"]:
        return (
            f"{month} 已有{basis_text} { _fmt(revenue) }。"
            "还没有支出，所以还不能判断这个月赚没有。导入支出后，把每一笔拖到对应的商品上。"
        )
    profit = summary["month_profit"]
    if summary["unknown_count"]:
        text = (
            f"{month} 按{basis_text}算，收入 { _fmt(revenue) }，目前导入的支出 { _fmt(summary['expense_total']) }。"
            f"如果这些就是全部支出，盈亏是 { _fmt(profit) }。"
            f"还有 {summary['unknown_count']} 个商品没挂成本，先不要把整月当成已经算清。"
        )
    else:
        verdict = "打平" if abs(profit) < 0.005 else ("赚钱" if profit > 0 else "亏损")
        shown = _fmt(0) if abs(profit) < 0.005 else _fmt(abs(profit))
        text = (
            f"{month} 按{basis_text}算，收入 { _fmt(revenue) }，支出 { _fmt(summary['expense_total']) }，"
            f"这个月{verdict} {shown}。"
        )
    if summary["basis"] == "confirmed" and summary["unconfirmed_amt"]:
        text += f" 另有未确认收货 { _fmt(summary['unconfirmed_amt']) }，没有算进这笔盈亏。"
    if summary["unallocated_total"] > 0:
        text += f" 其中 { _fmt(summary['unallocated_total']) } 支出还没挂到商品。整月盈亏已经扣掉了，但这笔成本还没算进任何一个商品。"
    if summary["pending_count"]:
        text += f" 另有 {summary['pending_count']} 笔退款没完结（{ _fmt(summary['pending_amt']) }），目前仍算在收入里。"
    return text


def _fmt(value):
    sign = "-" if value < 0 else ""
    return sign + f"¥{abs(value):,.2f}"


def load_demo(conn):
    month = "2099-01"
    if conn.execute("SELECT 1 FROM imports WHERE note='demo'").fetchone():
        return month
    now = datetime.now().isoformat(timespec="seconds")
    cur = conn.execute(
        "INSERT INTO imports (kind, filename, created_at, row_count, month_hint, note) VALUES (?,?,?,?,?,?)",
        ("orders", "示例", now, 4, month, "demo"),
    )
    import_id = cur.lastrowid
    samples = [
        ("示例商品甲", "颜色:红", 10, 1000, "没有申请退款", f"{month}-02", f"{month}-08"),
        ("示例商品甲", "颜色:蓝", 4, 400, "没有申请退款", f"{month}-03", ""),
        ("示例商品乙", "套餐:标准", 2, 500, "没有申请退款", f"{month}-04", f"{month}-09"),
        ("示例商品乙", "套餐:标准", 1, 80, "退款成功", f"{month}-05", f"{month}-10"),
    ]
    for index, (title, attr, qty, paid, status, pay, confirm) in enumerate(samples):
        conn.execute(
            """INSERT INTO revenue_lines (
                fingerprint, import_id, order_no, product_id, title, attr, package,
                qty, paid, refund_status, refund_amount, pay_time, pay_month,
                confirm_time, confirm_month, source_file
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                fingerprint(["demo", index]), import_id, f"D{index}", "", title, attr, "",
                qty, paid, status, paid if status == "退款成功" else 0,
                pay, month, confirm, month_of(confirm), "示例",
            ),
        )
    wood = conn.execute(
        """INSERT INTO expenses (fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (fingerprint(["demo", "wood"]), import_id, "示例木板", 300, f"{month}-06", month, "物料", "", "示例"),
    ).lastrowid
    conn.execute(
        "INSERT INTO allocations (expense_id, title, attr, amount) VALUES (?,?,?,?)",
        (wood, "示例商品甲", "", 300),
    )
    conn.execute(
        """INSERT INTO expenses (fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
           VALUES (?,?,?,?,?,?,?,?,?)""",
        (fingerprint(["demo", "pack"]), import_id, "示例包装", 120, f"{month}-07", month, "物流", "还没决定算到哪个商品", "示例"),
    )
    conn.commit()
    return month


def export_workbook(conn, month, basis="confirmed"):
    state = build_state(conn, month, basis)
    wb = openpyxl.Workbook()
    _write_summary_sheet(wb.active, state)
    _write_product_sheet(wb.create_sheet("按商品"), state)
    _write_attr_sheet(wb.create_sheet("按商品属性"), state)
    _write_expense_sheet(wb.create_sheet("支出"), state)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _header_fill():
    fill = PatternFill("solid", fgColor="0E5C56")
    font = Font(color="FFFFFF", bold=True, name="微软雅黑", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    return fill, font, align


def _write_header(ws, headers):
    fill, font, align = _header_fill()
    thin = Border(
        left=Side(style="thin", color="E2D5BE"),
        right=Side(style="thin", color="E2D5BE"),
        top=Side(style="thin", color="E2D5BE"),
        bottom=Side(style="thin", color="E2D5BE"),
    )
    for col, text in enumerate(headers, 1):
        cell = ws.cell(1, col, text)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = thin
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _write_summary_sheet(ws, state):
    ws.title = "本月结论"
    summary = state["summary"]
    ws["A1"] = "产品月账"
    ws["A1"].font = Font(name="微软雅黑", size=18, bold=True, color="221C14")
    ws["A3"] = summary["sentence"]
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:B3")
    ws.row_dimensions[3].height = 48
    rows = [
        ("月份", summary["month"]),
        ("收入口径", "只算确认收货" if summary["basis"] == "confirmed" else "确认收货 + 未确认"),
        ("确认收货数量", summary["confirmed_qty"]),
        ("确认收货金额", summary["confirmed_amt"]),
        ("未确认数量", summary["unconfirmed_qty"]),
        ("未确认金额", summary["unconfirmed_amt"]),
        ("退款成功（已剔除）", summary["refund_amt"]),
        ("支出合计", summary["expense_total"]),
        ("已挂到商品", summary["allocated_total"]),
        ("还没挂到商品", summary["unallocated_total"]),
        ("整月盈亏", summary["month_profit"] if summary["month_profit"] is not None else "还没有支出"),
        ("赚钱的商品", summary["profit_count"]),
        ("亏损的商品", summary["loss_count"]),
        ("还没挂成本的商品", summary["unknown_count"]),
    ]
    for index, (label, value) in enumerate(rows, 5):
        ws.cell(index, 1, label).font = Font(name="微软雅黑", bold=True)
        ws.cell(index, 2, value).font = Font(name="微软雅黑")
        if isinstance(value, float):
            ws.cell(index, 2).number_format = '#,##0.00'
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 88


def _write_product_sheet(ws, state):
    headers = ["商品标题", "赛道", "确认数量", "确认金额", "未确认数量", "未确认金额", "退款金额", "已挂成本", "利润", "盈亏", "成本明细"]
    _write_header(ws, headers)
    for row_index, prod in enumerate(state["products"], 2):
        detail = "；".join(f"{a['name']} {a['amount']:.2f}" for a in prod["allocations"])
        status = {"profit": "赚钱", "loss": "亏损", "even": "打平", "unknown": "成本未挂", "none": "无收入"}[prod["status"]]
        values = [
            prod["title"], prod["track"], prod["confirmed_qty"], prod["confirmed_amt"],
            prod["unconfirmed_qty"], prod["unconfirmed_amt"], prod["refund_amt"], prod["cost"],
            prod["profit"] if prod["profit"] is not None else "", status, detail,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="微软雅黑", size=10)
            if isinstance(value, float):
                cell.number_format = '#,##0.00'
        if prod["status"] == "loss":
            ws.cell(row_index, 9).font = Font(name="微软雅黑", color="A33B24", bold=True)
        elif prod["status"] == "profit":
            ws.cell(row_index, 9).font = Font(name="微软雅黑", color="2D6A3F", bold=True)
    widths = [42, 14, 12, 14, 12, 14, 12, 12, 12, 12, 48]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_attr_sheet(ws, state):
    headers = ["商品标题", "商品属性", "确认数量", "确认金额", "未确认数量", "未确认金额", "占该商品金额", "直接成本", "按金额估算成本", "估算利润"]
    _write_header(ws, headers)
    row_index = 2
    for prod in state["products"]:
        for attr in prod["attrs"]:
            values = [
                prod["title"], attr["attr"], attr["confirmed_qty"], attr["confirmed_amt"],
                attr["unconfirmed_qty"], attr["unconfirmed_amt"], attr["share"],
                attr["direct_cost"], attr["estimated_cost"] if attr["estimated_cost"] is not None else "",
                attr["estimated_profit"] if attr["estimated_profit"] is not None else "",
            ]
            for col, value in enumerate(values, 1):
                cell = ws.cell(row_index, col, value)
                cell.font = Font(name="微软雅黑", size=10)
                if isinstance(value, float):
                    cell.number_format = '0.00%' if col == 7 else '#,##0.00'
            row_index += 1
    widths = [42, 46, 12, 14, 12, 14, 14, 12, 18, 14]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _write_expense_sheet(ws, state):
    headers = ["项目", "金额", "还没挂出", "日期", "分类", "备注", "挂到的商品"]
    _write_header(ws, headers)
    for row_index, expense in enumerate(state["expenses"], 2):
        linked = "；".join(f"{a['title']} {a['amount']:.2f}" for a in expense["allocations"]) or "（还没挂）"
        values = [
            expense["name"], expense["amount"], expense["remaining"], expense["spent_on"],
            expense["category"], expense["note"], linked,
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row_index, col, value)
            cell.font = Font(name="微软雅黑", size=10)
            if isinstance(value, float):
                cell.number_format = '#,##0.00'
    widths = [28, 12, 12, 20, 12, 28, 48]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def expense_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "支出"
    headers = ["日期", "项目", "金额", "分类", "备注", "对应商品"]
    _write_header(ws, headers)
    ws.append(["2026-10-06", "PCB板加工", 1200, "加工", "10月这批板", "把商品标题写在这里就会自动挂上"])
    ws.append(["2026-10-08", "快递费", 86, "物流", "", ""])
    for col, width in enumerate([16, 24, 12, 12, 28, 42], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    note = wb.create_sheet("怎么填")
    note["A1"] = "每一行是一笔支出。金额填正数。"
    note["A2"] = "日期用来判断属于哪个月。没有日期时，导入软件里选的月份会用在所有空日期的行上。"
    note["A3"] = "对应商品可以空着。空着的支出导入后出现在右侧，用鼠标拖到左边的商品上。"
    note["A4"] = "对应商品如果写成完整的商品标题，或者写成型号（例如 YDI-D14），而且只对得上一个商品，会自动挂上。"
    note.column_dimensions["A"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def backup_data(conn):
    def rows(sql):
        return [dict(r) for r in conn.execute(sql)]
    return {
        "format": "product-ledger-backup",
        "version": 1,
        "exported_at": datetime.now().isoformat(timespec="seconds"),
        "imports": rows("SELECT * FROM imports"),
        "revenue_lines": rows("SELECT * FROM revenue_lines"),
        "expenses": rows("SELECT * FROM expenses"),
        "allocations": rows("SELECT * FROM allocations"),
        "product_meta": rows("SELECT * FROM product_meta"),
        "product_flags": rows("SELECT * FROM product_flags"),
    }


def restore_data(conn, payload):
    if not isinstance(payload, dict) or payload.get("format") != "product-ledger-backup":
        raise ValueError("这不是产品月账的备份文件。")
    conn.execute("DELETE FROM allocations")
    conn.execute("DELETE FROM expenses")
    conn.execute("DELETE FROM revenue_lines")
    conn.execute("DELETE FROM imports")
    conn.execute("DELETE FROM product_meta")
    conn.execute("DELETE FROM product_flags")
    for row in payload.get("imports") or []:
        conn.execute(
            "INSERT INTO imports (id, kind, filename, created_at, row_count, month_hint, note) VALUES (?,?,?,?,?,?,?)",
            (row["id"], row["kind"], row.get("filename"), row["created_at"], row.get("row_count") or 0, row.get("month_hint") or "", row.get("note") or ""),
        )
    for row in payload.get("revenue_lines") or []:
        conn.execute(
            """INSERT INTO revenue_lines (
                id, fingerprint, import_id, order_no, product_id, title, attr, package,
                qty, paid, refund_status, refund_amount, pay_time, pay_month,
                confirm_time, confirm_month, source_file
            ) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (
                row["id"], row["fingerprint"], row.get("import_id"), row.get("order_no"), row.get("product_id"),
                row["title"], row.get("attr") or "", row.get("package"), row.get("qty") or 0, row.get("paid") or 0,
                row.get("refund_status"), row.get("refund_amount") or 0, row.get("pay_time"), row.get("pay_month"),
                row.get("confirm_time"), row.get("confirm_month"), row.get("source_file"),
            ),
        )
    for row in payload.get("expenses") or []:
        conn.execute(
            """INSERT INTO expenses (id, fingerprint, import_id, name, amount, spent_on, month, category, note, source_file)
               VALUES (?,?,?,?,?,?,?,?,?,?)""",
            (
                row["id"], row.get("fingerprint"), row.get("import_id"), row["name"], row["amount"],
                row.get("spent_on"), row["month"], row.get("category") or "", row.get("note") or "", row.get("source_file"),
            ),
        )
    for row in payload.get("allocations") or []:
        conn.execute(
            "INSERT INTO allocations (id, expense_id, title, attr, amount) VALUES (?,?,?,?,?)",
            (row["id"], row["expense_id"], row["title"], row.get("attr") or "", row["amount"]),
        )
    for row in payload.get("product_meta") or []:
        conn.execute(
            "INSERT INTO product_meta (title, track, note) VALUES (?,?,?)",
            (row["title"], row.get("track") or "", row.get("note") or ""),
        )
    for row in payload.get("product_flags") or []:
        conn.execute(
            "INSERT INTO product_flags (title, month, no_cost) VALUES (?,?,?)",
            (row["title"], row["month"], row.get("no_cost") or 0),
        )
    conn.commit()


def sweep_tmp():
    ensure_dirs()
    cutoff = datetime.now().timestamp() - 86400
    for path in TMP.iterdir():
        try:
            if path.stat().st_mtime < cutoff:
                path.unlink()
        except OSError:
            pass
